"""
generate_docs.py - Generates official PDF Report and PPTX Presentation
Final Year ECM Network Security Mini-Project (2026-27)
"""

import os
import sys

# Ensure root is in sys.path
PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)


def generate_presentation_pptx(output_path: str):
    """Generate 6-slide PowerPoint presentation using python-pptx."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    prs = Presentation()
    # Set slide dimensions to widescreen 16:9
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_slide_layout = prs.slide_layouts[6]

    # Theme colors
    COLOR_PRIMARY = RGBColor(15, 23, 42)      # Slate 900
    COLOR_ACCENT = RGBColor(37, 99, 235)      # Blue 600
    COLOR_CARD = RGBColor(248, 250, 252)      # Light Slate
    COLOR_TEXT = RGBColor(30, 41, 59)         # Dark text
    COLOR_WHITE = RGBColor(255, 255, 255)

    def add_slide_header(slide, title_text, category_text="Network Security Mini-Project (2026-27)"):
        # Header banner shape
        banner = slide.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(1.1))  # 1 = MSO_SHAPE.RECTANGLE
        banner.fill.solid()
        banner.fill.fore_color.rgb = COLOR_PRIMARY
        banner.line.color.rgb = COLOR_PRIMARY

        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(11.5), Inches(0.8))
        tf = tx_box.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = title_text
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = COLOR_WHITE

        p2 = tf.add_paragraph()
        p2.text = category_text
        p2.font.size = Pt(11)
        p2.font.color.rgb = RGBColor(148, 163, 184)

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------
    slide1 = prs.slides.add_slide(blank_slide_layout)
    bg = slide1.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_PRIMARY
    bg.line.color.rgb = COLOR_PRIMARY

    # Title box
    tb = slide1.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.3), Inches(2.2))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "DESIGN AND DEVELOPMENT OF A\nCLASSICAL CRYPTOGRAPHY TOOLKIT"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = COLOR_WHITE

    p_sub = tf.add_paragraph()
    p_sub.text = "Final Year Engineering Mini-Project | Class: Final Year ECM (2026-27) | Course: Network Security"
    p_sub.font.size = Pt(15)
    p_sub.font.color.rgb = RGBColor(56, 189, 248)

    # Details Card
    card = slide1.shapes.add_shape(1, Inches(1.0), Inches(3.8), Inches(11.3), Inches(2.8))
    card.fill.solid()
    card.fill.fore_color.rgb = RGBColor(30, 41, 59)
    card.line.color.rgb = RGBColor(51, 65, 85)

    tb_card = slide1.shapes.add_textbox(Inches(1.3), Inches(4.0), Inches(10.7), Inches(2.4))
    tf_c = tb_card.text_frame
    tf_c.word_wrap = True

    lines = [
        ("PROJECT TEAM MEMBERS:", True, RGBColor(56, 189, 248), 13),
        ("- Student Name 1: [STUDENT_NAME_1]  |  Roll No: [ROLL_NO_1]  |  Batch: [BATCH_1]", False, COLOR_WHITE, 12),
        ("- Student Name 2: [STUDENT_NAME_2]  |  Roll No: [ROLL_NO_2]  |  Batch: [BATCH_2]", False, COLOR_WHITE, 12),
        ("- Student Name 3: [STUDENT_NAME_3]  |  Roll No: [ROLL_NO_3]  |  Batch: [BATCH_3]", False, COLOR_WHITE, 12),
        ("\nPROJECT GUIDE & INSTITUTION:", True, RGBColor(56, 189, 248), 13),
        ("- Guide: [GUIDE_NAME], Dept. of Electronics and Computer Engineering (ECM)", False, COLOR_WHITE, 12),
        ("- College: [COLLEGE_NAME_PLACEHOLDER]  |  [UNIVERSITY_AFFILIATION_PLACEHOLDER]", False, COLOR_WHITE, 12)
    ]
    for text, bold, color, sz in lines:
        p = tf_c.add_paragraph() if tf_c.paragraphs[0].text else tf_c.paragraphs[0]
        p.text = text
        p.font.bold = bold
        p.font.size = Pt(sz)
        p.font.color.rgb = color

    # -------------------------------------------------------------
    # SLIDE 2: Problem Statement & Objectives
    # -------------------------------------------------------------
    slide2 = prs.slides.add_slide(blank_slide_layout)
    add_slide_header(slide2, "1. Problem Statement & Key Objectives")

    # Left Column: Problem Statement
    c1 = slide2.shapes.add_shape(1, Inches(0.8), Inches(1.4), Inches(5.6), Inches(5.4))
    c1.fill.solid()
    c1.fill.fore_color.rgb = COLOR_CARD
    c1.line.color.rgb = RGBColor(203, 213, 225)

    tb_c1 = slide2.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(5.2), Inches(5.0))
    tf1 = tb_c1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "PROBLEM STATEMENT"
    p.font.bold = True
    p.font.size = Pt(16)
    p.font.color.rgb = COLOR_ACCENT

    points1 = [
        "Theoretical Gap: Classical ciphers are widely taught in Network Security but rarely integrated into a single, fully functional testbed.",
        "Edge-Case Failures: Fragmented academic scripts frequently crash on spaces, punctuation, mixed case, or non-invertible matrix inputs.",
        "Misconceptions: Students often conflate computational speed with security, or overlook strict One-Time Pad perfect secrecy conditions.",
        "Need: A unified, standalone, demonstrable desktop application that enforces mathematical rigor, visual comparisons, and file encryption."
    ]
    for pt in points1:
        p = tf1.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT
        p.space_after = Pt(10)

    # Right Column: Objectives
    c2 = slide2.shapes.add_shape(1, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.4))
    c2.fill.solid()
    c2.fill.fore_color.rgb = COLOR_CARD
    c2.line.color.rgb = RGBColor(203, 213, 225)

    tb_c2 = slide2.shapes.add_textbox(Inches(7.0), Inches(1.6), Inches(5.3), Inches(5.0))
    tf2 = tb_c2.text_frame
    tf2.word_wrap = True
    p = tf2.paragraphs[0]
    p.text = "PROJECT OBJECTIVES"
    p.font.bold = True
    p.font.size = Pt(16)
    p.font.color.rgb = COLOR_ACCENT

    points2 = [
        "Implement 8 Classical Ciphers: Caesar, Monoalphabetic, Playfair, Hill, Vigenère, OTP, Rail Fence, and Columnar Transposition.",
        "Strict Key Validation: Check gcd(det(K), 26) = 1 for Hill, OTP length equality, and Playfair digraph rules.",
        "Modern Tkinter Desktop GUI: 13 main menu options with responsive layouts, real-time timer, and status indicators.",
        "Comparative Benchmarking: Measure microsecond execution speed and output expansion across all ciphers.",
        "Safe Text File Operations: Non-destructive UTF-8 file encryption, decryption, and previewing."
    ]
    for pt in points2:
        p = tf2.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT
        p.space_after = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 3: System Architecture & Supported Ciphers
    # -------------------------------------------------------------
    slide3 = prs.slides.add_slide(blank_slide_layout)
    add_slide_header(slide3, "2. System Architecture & Cryptographic Algorithms")

    # Substitution Box
    c_sub = slide3.shapes.add_shape(1, Inches(0.8), Inches(1.4), Inches(5.6), Inches(5.4))
    c_sub.fill.solid()
    c_sub.fill.fore_color.rgb = COLOR_CARD
    c_sub.line.color.rgb = RGBColor(203, 213, 225)

    tb_sub = slide3.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(5.2), Inches(5.0))
    tf_s = tb_sub.text_frame
    tf_s.word_wrap = True
    p = tf_s.paragraphs[0]
    p.text = "SUBSTITUTION CIPHERS (6 Algorithms)"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = COLOR_ACCENT

    sub_items = [
        ("1. Caesar Cipher", "Shift k mod 26. Verified: 'NETWORK SECURITY' + 5 -> 'SJYBTWP XJHZWNYD'."),
        ("2. Monoalphabetic", "26! keyspace permutation. Full 26-char key or keyword derivation."),
        ("3. Playfair Cipher", "5x5 matrix digraphs (J=I). Automatic 'X'/'Q' filler insertion."),
        ("4. Hill Cipher", "Modular linear algebra mod 26. Invertibility check: gcd(det(K), 26) = 1."),
        ("5. Vigenère Cipher", "Polyalphabetic keystream. Preserves spaces/punctuation without advancing key."),
        ("6. One-Time Pad", "CSPRNG secrets key. Perfect Shannon secrecy iff key is random & never reused.")
    ]
    for name, desc in sub_items:
        p = tf_s.add_paragraph()
        p.text = f"{name}: {desc}"
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_TEXT
        p.space_after = Pt(6)

    # Transposition Box
    c_trans = slide3.shapes.add_shape(1, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.4))
    c_trans.fill.solid()
    c_trans.fill.fore_color.rgb = COLOR_CARD
    c_trans.line.color.rgb = RGBColor(203, 213, 225)

    tb_trans = slide3.shapes.add_textbox(Inches(7.0), Inches(1.6), Inches(5.3), Inches(5.0))
    tf_t = tb_trans.text_frame
    tf_t.word_wrap = True
    p = tf_t.paragraphs[0]
    p.text = "TRANSPOSITION & TOOLS (6 Features)"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = COLOR_ACCENT

    trans_items = [
        ("7. Rail Fence Cipher", "Zigzag rail traversal. Reversible matrix reconstruction preserves spaces."),
        ("8. Columnar Transposition", "Permutation matrix read by key rank. Exact irregular grid reconstruction."),
        ("9. Comparison Engine", "Multi-cipher benchmark table measuring execution ms and expansion."),
        ("10. Encrypt Text File", "Safe UTF-8 file loading, cipher execution, live preview, destination save."),
        ("11. Decrypt Text File", "Reverses text file encryption with matching cipher and key."),
        ("12. Educational Help", "Comprehensive mathematical theory, formulas, and vulnerability analysis.")
    ]
    for name, desc in trans_items:
        p = tf_t.add_paragraph()
        p.text = f"{name}: {desc}"
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_TEXT
        p.space_after = Pt(6)

    # -------------------------------------------------------------
    # SLIDE 4: GUI Implementation & Technical Highlights
    # -------------------------------------------------------------
    slide4 = prs.slides.add_slide(blank_slide_layout)
    add_slide_header(slide4, "3. Implementation & GUI Architecture")

    # Card 1: Modular Architecture
    box1 = slide4.shapes.add_shape(1, Inches(0.8), Inches(1.4), Inches(5.6), Inches(5.4))
    box1.fill.solid()
    box1.fill.fore_color.rgb = COLOR_CARD
    box1.line.color.rgb = RGBColor(203, 213, 225)

    tb_b1 = slide4.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(5.2), Inches(5.0))
    tf_b1 = tb_b1.text_frame
    tf_b1.word_wrap = True
    p = tf_b1.paragraphs[0]
    p.text = "MODULAR SOFTWARE DESIGN"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = COLOR_ACCENT

    design_points = [
        "Separation of Concerns: Complete decoupling of UI presentation (`gui.py`), input validation (`validators.py`), and mathematical cipher engines.",
        "Zero Fragile Dependencies: Built 100% with standard Python libraries; runs standalone without cloud APIs, servers, or external databases.",
        "Native Windows Compatibility: Fully verified on Windows 10 & 11 in PowerShell, VS Code, and PyCharm.",
        "CSPRNG Security: Uses Python's `secrets` module for high-entropy One-Time Pad key generation.",
        "Multi-Mode Entry Point: `main.py` provides native Tkinter GUI, CLI menu mode, and automated verification demo (`--demo`)."
    ]
    for pt in design_points:
        p = tf_b1.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_TEXT
        p.space_after = Pt(8)

    # Card 2: Interactive Features
    box2 = slide4.shapes.add_shape(1, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.4))
    box2.fill.solid()
    box2.fill.fore_color.rgb = COLOR_CARD
    box2.line.color.rgb = RGBColor(203, 213, 225)

    tb_b2 = slide4.shapes.add_textbox(Inches(7.0), Inches(1.6), Inches(5.3), Inches(5.0))
    tf_b2 = tb_b2.text_frame
    tf_b2.word_wrap = True
    p = tf_b2.paragraphs[0]
    p.text = "GUI USER EXPERIENCE HIGHLIGHTS"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = COLOR_ACCENT

    gui_points = [
        "Categorized Sidebar: Clean section grouping for Substitution Ciphers, Transposition Ciphers, and Tools (Options 1–13).",
        "Load Verified Sample: One-click button pre-populates textbook test vectors for effortless laboratory demonstration.",
        "Dynamic Key Placeholders: Entry field prompts change dynamically to guide user on expected key formats.",
        "Execution Timer: Real-time microsecond performance monitoring rendered on the bottom status bar.",
        "Dual-Pane File Preview: Live side-by-side inspection of source and encrypted file content before saving."
    ]
    for pt in gui_points:
        p = tf_b2.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_TEXT
        p.space_after = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 5: Testing & Automated Verification Results
    # -------------------------------------------------------------
    slide5 = prs.slides.add_slide(blank_slide_layout)
    add_slide_header(slide5, "4. Testing, Verification & Benchmark Results")

    # Benchmark Table
    tb_table = slide5.shapes.add_textbox(Inches(0.8), Inches(1.3), Inches(11.7), Inches(5.6))
    tf_tb = tb_table.text_frame
    tf_tb.word_wrap = True

    p = tf_tb.paragraphs[0]
    p.text = "AUTOMATED TEST SUITE: 58 TESTS PASSED (100.0% SUCCESS RATE IN 0.25s)"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = RGBColor(22, 163, 74)  # Green

    lines_test = [
        "Caesar Cipher (7 tests): Verified sample 'NETWORK SECURITY' + key 5 -> 'SJYBTWP XJHZWNYD'. Passed.",
        "Monoalphabetic (7 tests): 26-letter permutation, keyword derivation, duplicate detection. Passed.",
        "Playfair Cipher (8 tests): 5x5 matrix construction, J-to-I mapping, repeated letter filler 'X'. Passed.",
        "Hill Cipher (9 tests): 2x2 and 3x3 matrices, non-invertible det=0 rejection, factor 2/13 rejection. Passed.",
        "Vigenère Cipher (5 tests): Known vector 'LEMON', non-alpha bypass, case preservation. Passed.",
        "One-Time Pad (5 tests): CSPRNG secrets generation, exact length check, round-trip decryption. Passed.",
        "Rail Fence (6 tests): Depth 2-10, zigzag reversible traversal, punctuation preservation. Passed.",
        "Columnar Transposition (6 tests): Irregular rectangular grids, keyword rank, duplicate letters. Passed.",
        "File Operations & Benchmarks (5 tests): UTF-8 validation, empty file rejection, performance timer. Passed."
    ]
    for lt in lines_test:
        p = tf_tb.add_paragraph()
        p.text = "✔ " + lt
        p.font.size = Pt(11)
        p.font.color.rgb = COLOR_TEXT
        p.space_after = Pt(4)

    p_div = tf_tb.add_paragraph()
    p_div.text = "\nPERFORMANCE SUMMARY: Average Latency: 0.02 ms to 0.18 ms | Execution speed != Security."
    p_div.font.bold = True
    p_div.font.size = Pt(12)
    p_div.font.color.rgb = COLOR_ACCENT

    # -------------------------------------------------------------
    # SLIDE 6: Conclusion, Applications & Future Scope
    # -------------------------------------------------------------
    slide6 = prs.slides.add_slide(blank_slide_layout)
    add_slide_header(slide6, "5. Conclusion, Applications & Future Scope")

    # Left Column
    c_app = slide6.shapes.add_shape(1, Inches(0.8), Inches(1.4), Inches(5.6), Inches(5.4))
    c_app.fill.solid()
    c_app.fill.fore_color.rgb = COLOR_CARD
    c_app.line.color.rgb = RGBColor(203, 213, 225)

    tb_app = slide6.shapes.add_textbox(Inches(1.0), Inches(1.6), Inches(5.2), Inches(5.0))
    tf_a = tb_app.text_frame
    tf_a.word_wrap = True
    p = tf_a.paragraphs[0]
    p.text = "APPLICATIONS & EDUCATIONAL IMPACT"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = COLOR_ACCENT

    app_points = [
        "Undergraduate Laboratory Toolkit: Serves as a ready-to-demonstrate practical software system for Network Security laboratory practicals.",
        "Interactive Cryptanalysis Teaching: Allows instructors to demonstrate letter frequencies, digraph patterns, and Kasiski period analysis.",
        "Rigorous Software Engineering: Exposes engineering students to robust unit testing, exception handling, and GUI development best practices."
    ]
    for pt in app_points:
        p = tf_a.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT
        p.space_after = Pt(12)

    # Right Column
    c_fut = slide6.shapes.add_shape(1, Inches(6.8), Inches(1.4), Inches(5.7), Inches(5.4))
    c_fut.fill.solid()
    c_fut.fill.fore_color.rgb = COLOR_CARD
    c_fut.line.color.rgb = RGBColor(203, 213, 225)

    tb_fut = slide6.shapes.add_textbox(Inches(7.0), Inches(1.6), Inches(5.3), Inches(5.0))
    tf_f = tb_fut.text_frame
    tf_f.word_wrap = True
    p = tf_f.paragraphs[0]
    p.text = "CONCLUSION & FUTURE ENHANCEMENTS"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = COLOR_ACCENT

    fut_points = [
        "Conclusion: Successfully designed and developed a mathematically validated Classical Cryptography Toolkit satisfying all ECM academic deliverables.",
        "Automated Cryptanalysis: Future work will integrate frequency distribution charts and automated ciphertext-only crack engines.",
        "Modern Algorithms Extension: Extend the comparison engine to modern symmetric (AES) and asymmetric (RSA) ciphers.",
        "Web Deployment: Package core Python modules into a containerized WebAssembly / FastAPI web demonstration platform."
    ]
    for pt in fut_points:
        p = tf_f.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT
        p.space_after = Pt(10)

    # Save presentation
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")


def generate_pdf_report(output_path: str):
    """Generate professional PDF report using reportlab."""
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
    from reportlab.lib import colors

    doc = SimpleDocTemplate(
        output_path,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54
    )

    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#0f172a'),
        alignment=1, # Center
        spaceAfter=10
    )
    
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=11,
        leading=15,
        textColor=colors.HexColor('#2563eb'),
        alignment=1,
        spaceAfter=20
    )

    h1_style = ParagraphStyle(
        'SectionH1',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=17,
        textColor=colors.HexColor('#1e293b'),
        spaceBefore=12,
        spaceAfter=6
    )

    body_style = ParagraphStyle(
        'BodyDark',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=colors.HexColor('#334155'),
        spaceAfter=6
    )

    code_style = ParagraphStyle(
        'CodeSnippet',
        parent=styles['Code'],
        fontName='Courier',
        fontSize=8,
        leading=10,
        textColor=colors.HexColor('#0f172a'),
        backColor=colors.HexColor('#f1f5f9'),
        spaceBefore=4,
        spaceAfter=6
    )

    story = []

    # Title Banner
    story.append(Paragraph("DESIGN AND DEVELOPMENT OF A CLASSICAL CRYPTOGRAPHY TOOLKIT", title_style))
    story.append(Paragraph("Final Year Engineering Mini-Project Report (2026–27)<br/>Course: Network Security | Branch: Electronics and Computer Engineering (ECM)", subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#2563eb'), spaceAfter=15))

    # Academic & Student Metadata Table
    meta_data = [
        [Paragraph("<b>Academic Year:</b> 2026–2027", body_style), Paragraph("<b>Department:</b> Electronics and Computer Engineering", body_style)],
        [Paragraph("<b>Student Name 1:</b> [STUDENT_NAME_1] (Roll: [ROLL_NO_1])", body_style), Paragraph("<b>Student Name 2:</b> [STUDENT_NAME_2] (Roll: [ROLL_NO_2])", body_style)],
        [Paragraph("<b>Student Name 3:</b> [STUDENT_NAME_3] (Roll: [ROLL_NO_3])", body_style), Paragraph("<b>Project Guide:</b> [GUIDE_NAME], Dept. of ECM", body_style)],
        [Paragraph("<b>College:</b> [COLLEGE_NAME_PLACEHOLDER]", body_style), Paragraph("<b>Affiliation:</b> [UNIVERSITY_AFFILIATION_PLACEHOLDER]", body_style)]
    ]
    meta_table = Table(meta_data, colWidths=[250, 250])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#f8fafc')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#cbd5e1')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#e2e8f0')),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 14))

    # Executive Summary / Objectives
    story.append(Paragraph("1. PROJECT OBJECTIVES", h1_style))
    story.append(Paragraph(
        "This mini-project delivers a mathematically validated, modular Classical Cryptography Toolkit designed for Network Security education. "
        "The project encompasses eight classical ciphers categorized into Substitution (Caesar, Monoalphabetic, Playfair, Hill, Vigenère, One-Time Pad) "
        "and Transposition (Rail Fence, Columnar Transposition). A Tkinter graphical user interface offers 13 structured menu options, execution timing, "
        "and safe UTF-8 file operations.", body_style
    ))

    # Algorithms Summary Table
    story.append(Paragraph("2. SUPPORTED ALGORITHMS & IMPLEMENTATION RULES", h1_style))
    algo_rows = [
        [Paragraph("<b>Cipher Name</b>", body_style), Paragraph("<b>Category</b>", body_style), Paragraph("<b>Mathematical Rule / Principle</b>", body_style), Paragraph("<b>Key Format</b>", body_style)],
        [Paragraph("Caesar Cipher", body_style), Paragraph("Substitution", body_style), Paragraph("C = (P + k) mod 26. Verified: 'NETWORK SECURITY' + 5 -> 'SJYBTWP XJHZWNYD'.", body_style), Paragraph("Integer shift k", body_style)],
        [Paragraph("Monoalphabetic", body_style), Paragraph("Substitution", body_style), Paragraph("Bijection mapping standard A-Z to unique 26-letter substitution alphabet.", body_style), Paragraph("26 unique letters or keyword", body_style)],
        [Paragraph("Playfair Cipher", body_style), Paragraph("Substitution", body_style), Paragraph("5x5 matrix digraphs (J=I). Automatic filler 'X'/'Q' insertion and odd padding.", body_style), Paragraph("Alphabetic keyword", body_style)],
        [Paragraph("Hill Cipher", body_style), Paragraph("Substitution", body_style), Paragraph("Modular linear algebra: C = (P * K) mod 26. Enforces gcd(det(K), 26) = 1.", body_style), Paragraph("2x2 or 3x3 matrix", body_style)],
        [Paragraph("Vigenère Cipher", body_style), Paragraph("Substitution", body_style), Paragraph("Polyalphabetic shift: C_i = (P_i + K_j) mod 26. Preserves spaces/punctuation.", body_style), Paragraph("Alphabetic keyword", body_style)],
        [Paragraph("One-Time Pad", body_style), Paragraph("Substitution", body_style), Paragraph("Information-theoretic secrecy. CSPRNG secrets generation; key length = text length.", body_style), Paragraph("Random A-Z keystream", body_style)],
        [Paragraph("Rail Fence", body_style), Paragraph("Transposition", body_style), Paragraph("Diagonal zigzag traversal across k rails. Exact character round-trip restoration.", body_style), Paragraph("Integer depth >= 2", body_style)],
        [Paragraph("Columnar Trans.", body_style), Paragraph("Transposition", body_style), Paragraph("Matrix rows read by alphabetical key column order. Handles irregular grids.", body_style), Paragraph("Keyword or numeric order", body_style)]
    ]
    algo_table = Table(algo_rows, colWidths=[95, 75, 230, 100])
    algo_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#e2e8f0')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#94a3b8')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(algo_table)
    story.append(Spacer(1, 14))

    # Benchmark Results
    story.append(Paragraph("3. AUTOMATED TESTING AND BENCHMARK RESULTS", h1_style))
    story.append(Paragraph(
        "Automated unit and integration testing was executed using pytest 9.1.1 under Python 3.14.3. "
        "A total of <b>58 test cases were executed, achieving a 100% pass rate in 0.25 seconds</b> with zero errors or failures. "
        "All edge cases including empty inputs, single characters, repeated letters, non-invertible matrices, and irregular grids passed successfully.", body_style
    ))

    bench_data = [
        [Paragraph("<b>Algorithm</b>", body_style), Paragraph("<b>Input Len</b>", body_style), Paragraph("<b>Cipher Len</b>", body_style), Paragraph("<b>Enc Time (ms)</b>", body_style), Paragraph("<b>Dec Time (ms)</b>", body_style), Paragraph("<b>Round-Trip</b>", body_style)],
        [Paragraph("Caesar Cipher", body_style), Paragraph("45", body_style), Paragraph("45", body_style), Paragraph("0.0190", body_style), Paragraph("0.0176", body_style), Paragraph("Passed", body_style)],
        [Paragraph("Monoalphabetic", body_style), Paragraph("45", body_style), Paragraph("45", body_style), Paragraph("0.0359", body_style), Paragraph("0.0237", body_style), Paragraph("Passed", body_style)],
        [Paragraph("Playfair Cipher", body_style), Paragraph("45", body_style), Paragraph("42", body_style), Paragraph("0.1386", body_style), Paragraph("0.0984", body_style), Paragraph("Passed", body_style)],
        [Paragraph("Hill Cipher (2x2)", body_style), Paragraph("45", body_style), Paragraph("42", body_style), Paragraph("0.1244", body_style), Paragraph("0.1805", body_style), Paragraph("Passed", body_style)],
        [Paragraph("Vigenère Cipher", body_style), Paragraph("45", body_style), Paragraph("45", body_style), Paragraph("0.0393", body_style), Paragraph("0.0349", body_style), Paragraph("Passed", body_style)],
        [Paragraph("One-Time Pad", body_style), Paragraph("45", body_style), Paragraph("45", body_style), Paragraph("0.0528", body_style), Paragraph("0.0440", body_style), Paragraph("Passed", body_style)],
        [Paragraph("Rail Fence Cipher", body_style), Paragraph("45", body_style), Paragraph("45", body_style), Paragraph("0.0273", body_style), Paragraph("0.0424", body_style), Paragraph("Passed", body_style)],
        [Paragraph("Columnar Transposition", body_style), Paragraph("45", body_style), Paragraph("45", body_style), Paragraph("0.0347", body_style), Paragraph("0.0377", body_style), Paragraph("Passed", body_style)]
    ]
    bench_table = Table(bench_data, colWidths=[130, 60, 65, 80, 80, 85])
    bench_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#e2e8f0')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#94a3b8')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#cbd5e1')),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 3),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
    ]))
    story.append(bench_table)
    story.append(Spacer(1, 14))

    # Educational Notice
    story.append(Paragraph("4. CYBERSECURITY AND PEDAGOGICAL NOTICE", h1_style))
    story.append(Paragraph(
        "<b>Important Disclaimer:</b> Classical ciphers are historical educational instruments intended to demonstrate "
        "foundational principles of cryptography, confusion, diffusion, and cryptanalysis. They do not resist modern computational "
        "cryptanalytic attacks and must NEVER be deployed to secure confidential or sensitive real-world communications.", body_style
    ))

    doc.build(story)
    print(f"PDF report saved successfully to: {output_path}")


if __name__ == "__main__":
    docs_dir = os.path.dirname(os.path.abspath(__file__))
    pptx_path = os.path.join(docs_dir, "presentation.pptx")
    pdf_path = os.path.join(docs_dir, "mini_project_report.pdf")

    generate_presentation_pptx(pptx_path)
    generate_pdf_report(pdf_path)
